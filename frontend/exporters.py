"""Create local DOCX and PDF exports for the offline demo mode."""

from __future__ import annotations

import io
from html import escape
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def make_docx(result: dict[str, Any]) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    title = document.add_heading(result.get("title") or "Протокол совещания", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = RGBColor(20, 47, 79)

    document.add_paragraph("Протокол совещания · локальная система автопротоколирования")
    document.add_heading("Краткое саммари", level=1)
    for item in result.get("summary", []):
        document.add_paragraph(str(item), style="List Bullet")

    document.add_heading("Поручения", level=1)
    tasks = result.get("tasks", [])
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, ["Поручение", "Ответственный", "Срок", "Статус"]):
        cell.text = label
    for task in tasks:
        cells = table.add_row().cells
        for cell, value in zip(
            cells,
            [task.get("description", ""), task.get("assignee", ""), task.get("deadline", ""), task.get("status", "В работе")],
        ):
            cell.text = str(value)

    document.add_heading("Транскрипт", level=1)
    for line in result.get("transcript", []):
        paragraph = document.add_paragraph()
        speaker = paragraph.add_run(_speaker_name(line.get("speaker", "")) + "  ")
        speaker.bold = True
        paragraph.add_run(str(line.get("text", "")))

    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _speaker_name(speaker: str) -> str:
    if speaker.startswith("SPEAKER_"):
        suffix = speaker.removeprefix("SPEAKER_")
        try:
            return f"Спикер {int(suffix) + 1}"
        except ValueError:
            pass
    return speaker or "Спикер"


def _register_local_font() -> tuple[str, str]:
    candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/segoeui.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
    ]
    bold_candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf"),
        Path("C:/Windows/Fonts/segoeuib.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
    ]
    regular = next((path for path in candidates if path.is_file()), None)
    bold = next((path for path in bold_candidates if path.is_file()), regular)
    if regular is None:
        raise RuntimeError("Не найден локальный шрифт с кириллицей для PDF")
    pdfmetrics.registerFont(TTFont("ProtocolSans", str(regular)))
    pdfmetrics.registerFont(TTFont("ProtocolSans-Bold", str(bold)))
    pdfmetrics.registerFontFamily("ProtocolSans", normal="ProtocolSans", bold="ProtocolSans-Bold",
                                 italic="ProtocolSans", boldItalic="ProtocolSans-Bold")
    return "ProtocolSans", "ProtocolSans-Bold"


def make_pdf(result: dict[str, Any]) -> bytes:
    regular_font, bold_font = _register_local_font()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ProtocolTitle", parent=styles["Title"], fontName=bold_font,
        fontSize=20, leading=25, textColor=colors.HexColor("#142f4f"),
        alignment=TA_LEFT, spaceAfter=8,
    )
    heading_style = ParagraphStyle(
        "ProtocolHeading", parent=styles["Heading2"], fontName=bold_font,
        fontSize=12, leading=16, textColor=colors.HexColor("#176b66"),
        spaceBefore=12, spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "ProtocolBody", parent=styles["BodyText"], fontName=regular_font,
        fontSize=9.5, leading=14, textColor=colors.HexColor("#223047"),
        spaceAfter=5,
    )
    small_style = ParagraphStyle(
        "ProtocolSmall", parent=body_style, fontSize=8, leading=11,
    )
    story = [
        Paragraph(escape(result.get("title") or "Протокол совещания"), title_style),
        Paragraph("Протокол совещания · локальная система автопротоколирования", small_style),
        Spacer(1, 8),
        Paragraph("Краткое саммари", heading_style),
    ]
    for item in result.get("summary", []):
        story.append(Paragraph("• " + escape(str(item)), body_style))

    story.append(Paragraph("Поручения", heading_style))
    table_data = [[
        Paragraph(label, small_style)
        for label in ("Поручение", "Ответственный", "Срок", "Статус")
    ]]
    for task in result.get("tasks", []):
        table_data.append([
            Paragraph(escape(str(task.get("description", ""))), small_style),
            Paragraph(escape(str(task.get("assignee", ""))), small_style),
            Paragraph(escape(str(task.get("deadline", ""))), small_style),
            Paragraph(escape(str(task.get("status", "В работе"))), small_style),
        ])
    if len(table_data) == 1:
        table_data.append([Paragraph("Поручения не обнаружены", small_style), "", "", ""])
    task_table = Table(table_data, colWidths=[230, 95, 82, 70], repeatRows=1, hAlign="LEFT")
    task_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf1f5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#142f4f")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d4dee7")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([task_table, Paragraph("Транскрипт", heading_style)])
    for line in result.get("transcript", []):
        label = escape(_speaker_name(line.get("speaker", "")))
        text = escape(str(line.get("text", "")))
        story.append(Paragraph(f"<b>{label}</b> — {text}", body_style))
    if result.get("warnings"):
        story.append(Paragraph("Примечания", heading_style))
        for warning in result["warnings"]:
            story.append(Paragraph("• " + escape(str(warning)), small_style))

    stream = io.BytesIO()
    SimpleDocTemplate(
        stream, pagesize=A4, rightMargin=42, leftMargin=42,
        topMargin=42, bottomMargin=42,
        title=result.get("title") or "Протокол совещания",
    ).build(story)
    return stream.getvalue()
