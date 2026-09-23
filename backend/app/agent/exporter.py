from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from .models import ProtocolResult

def export_files(result: ProtocolResult, folder: Path) -> tuple[Path, Path]:
    folder.mkdir(parents=True, exist_ok=True)
    docx_path, pdf_path = folder / "protocol.docx", folder / "protocol.pdf"
    doc = Document(); title = doc.add_heading("Протокол совещания", 0); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"ID: {result.meeting_id}")
    doc.add_heading("Саммари", level=1)
    for item in result.summary: doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Поручения", level=1)
    table = doc.add_table(rows=1, cols=4); table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Ответственный", "Срок", "Суть", "Цитата"]): cell.text = text
    for task in result.tasks:
        cells = table.add_row().cells
        for cell, text in zip(cells, [task.assignee, task.deadline, task.description, task.source_quote]): cell.text = text
    doc.add_heading("Транскрипт", level=1)
    for line in result.transcript: doc.add_paragraph(f"[{line.speaker}] {line.text}")
    doc.save(docx_path)
    pdf = canvas.Canvas(str(pdf_path), pagesize=A4); width, height = A4; y = height - 50
    pdf.setFont("Helvetica-Bold", 16); pdf.drawString(40, y, "Meeting protocol"); y -= 30; pdf.setFont("Helvetica", 10)
    for section in ["Summary: "+"; ".join(result.summary), "Tasks: "+str(len(result.tasks))] + [f"[{x.speaker}] {x.text}" for x in result.transcript]:
        for chunk in [section[i:i+100] for i in range(0, len(section), 100)]:
            if y < 45: pdf.showPage(); y = height - 50; pdf.setFont("Helvetica", 10)
            pdf.drawString(40, y, chunk); y -= 15
    pdf.save(); return docx_path, pdf_path

