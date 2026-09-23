"""Local Unicode DOCX/PDF protocol export; no network or model dependencies."""

from __future__ import annotations

import hashlib
import math
import os
import tempfile
import threading
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape


_FONT_LOCK = threading.Lock()


def font_path() -> Path:
    """Find a local Unicode TTF, or fail explicitly before accepting export work."""
    configured = os.environ.get("PROTOCOL_FONT_PATH")
    if configured:
        selected = Path(configured).expanduser()
        if selected.is_file():
            return selected
        raise RuntimeError("Настроенный Unicode-шрифт для PDF недоступен.")

    windows = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    candidates = (
        windows / "arial.ttf",
        windows / "segoeui.ttf",
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/dejavu-sans-fonts/DejaVuSans.ttf"),
        Path("/usr/share/fonts/TTF/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
        Path("/Library/Fonts/Arial.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
    )
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise RuntimeError(
        "Для локального PDF нужен установленный Arial или DejaVu Sans. "
        "Укажите локальный TTF-файл в PROTOCOL_FONT_PATH."
    )


def _clean(value: Any, fallback: str = "не указано") -> str:
    text = fallback if value is None or not str(value).strip() else str(value)
    # Both OOXML and ReportLab's paragraph parser require XML 1.0 characters.
    text = "".join(
        character
        if (
            character in "\t\n\r"
            or 0x20 <= ord(character) <= 0xD7FF
            or 0xE000 <= ord(character) <= 0xFFFD
            or 0x10000 <= ord(character) <= 0x10FFFF
        )
        else "\ufffd"
        for character in text
    )
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _timestamp(value: Any) -> str:
    try:
        seconds = float(value)
        if not math.isfinite(seconds) or seconds < 0:
            return "неизвестное время"
        milliseconds = round(seconds * 1000)
    except (TypeError, ValueError, OverflowError):
        return "неизвестное время"
    hours, remainder = divmod(milliseconds, 3_600_000)
    minutes, remainder = divmod(remainder, 60_000)
    seconds, fraction = divmod(remainder, 1000)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}.{fraction:03d}"


def _paragraphs(title: str, result: dict) -> list[tuple[str, str]]:
    paragraphs = [("title", _clean(title, "Протокол совещания"))]
    if result.get("is_demo") is True:
        paragraphs.append(("body", "DEMO — синтетический пример, не результат обработки реального совещания."))
    paragraphs.append(
        (
            "subheading",
            "Автоматический черновик. Проверьте по записи; метки говорящих не являются именами.",
        )
    )
    paragraphs.extend(
        [
            ("heading", "Резюме"),
            ("body", _clean(result.get("summary"), "Резюме не предоставлено.")),
            ("heading", "Поручения"),
        ]
    )
    tasks = result.get("tasks") or []
    if not tasks:
        paragraphs.append(("body", "Поручения не предоставлены."))
    for index, task in enumerate(tasks, 1):
        paragraphs.append(("subheading", f"{index}. {_clean(task.get('task'), 'Текст поручения не указан')}"))
        for key, label in (
            ("assignee", "Ответственный"),
            ("deadline", "Срок"),
            ("assigned_by", "Поручил"),
            ("source_speaker", "Говорящий в источнике"),
            ("source_text", "Исходная реплика"),
        ):
            paragraphs.append(("body", f"{label}: {_clean(task.get(key))}"))

    paragraphs.append(("heading", "Транскрипт"))
    transcript = result.get("transcript") or []
    if not transcript:
        paragraphs.append(("body", "Реплики не предоставлены."))
    for segment in transcript:
        caption = (
            f"{_timestamp(segment.get('start'))} — {_timestamp(segment.get('end'))}"
            f"  ·  {_clean(segment.get('speaker'), 'Говорящий не определён')}"
        )
        languages = segment.get("detected_languages") or []
        if languages:
            caption += "  ·  Языки: " + ", ".join(_clean(code) for code in languages)
        paragraphs.append(("subheading", caption))
        paragraphs.append(("body", _clean(segment.get("text"), "Текст реплики отсутствует.")))
    return paragraphs


def _pdf_fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    selected = font_path()
    font_name = "SamrukUnicode" + hashlib.sha256(str(selected.resolve()).encode("utf-8")).hexdigest()[:12]
    bold_name = font_name + "Bold"
    bold_files = {
        "arial.ttf": "arialbd.ttf",
        "segoeui.ttf": "segoeuib.ttf",
        "DejaVuSans.ttf": "DejaVuSans-Bold.ttf",
        "LiberationSans-Regular.ttf": "LiberationSans-Bold.ttf",
        "Arial.ttf": "Arial Bold.ttf",
    }
    bold_path = selected.with_name(bold_files.get(selected.name, selected.name))
    if not bold_path.is_file():
        bold_path = selected
    with _FONT_LOCK:
        if font_name not in pdfmetrics.getRegisteredFontNames():
            regular = TTFont(font_name, str(selected))
            required = "АяӘәҒғҚқҢңӨөҰұҮүҺһІі"
            if any(ord(character) not in regular.face.charToGlyph for character in required):
                raise RuntimeError("Шрифт PDF не содержит необходимые русские и казахские символы.")
            pdfmetrics.registerFont(regular)
            pdfmetrics.registerFont(TTFont(bold_name, str(bold_path)))
            pdfmetrics.registerFontFamily(font_name, normal=font_name, bold=bold_name, italic=font_name, boldItalic=bold_name)
    return font_name, bold_name


def _write_docx(path: Path, title: str, paragraphs: list[tuple[str, str]]) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    document = Document()
    document.core_properties.title = _clean(title, "Протокол совещания")
    document.core_properties.author = "SAMRUK KAZYNA"
    document.core_properties.last_modified_by = "SAMRUK KAZYNA"
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
        style.element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Arial")
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    document.sections[0].header.paragraphs[0].text = "SAMRUK KAZYNA · Протокол совещания"
    footer = section.footer.paragraphs[0]
    footer.add_run("Страница ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    styles = {"title": "Title", "heading": "Heading 1", "subheading": "Heading 2", "body": "Normal"}
    for kind, text in paragraphs:
        # python-docx escapes XML and preserves Unicode/newlines itself.
        paragraph = document.add_paragraph(text, style=styles[kind])
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.keep_with_next = kind in ("title", "heading", "subheading")
    document.save(str(path))


def _write_pdf(path: Path, title: str, paragraphs: list[tuple[str, str]], fonts: tuple[str, str]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    regular, bold = fonts
    styles = {}
    for kind, size, leading, after in (
        ("title", 20, 25, 16),
        ("heading", 14, 18, 9),
        ("subheading", 11, 15, 5),
        ("body", 10, 14, 7),
    ):
        styles[kind] = ParagraphStyle(
            name=kind,
            fontName=regular if kind == "body" else bold,
            fontSize=size,
            leading=leading,
            spaceAfter=after,
            spaceBefore=8 if kind == "heading" else 0,
            alignment=TA_LEFT,
            textColor=colors.HexColor("#17212C"),
            splitLongWords=True,
            allowWidows=True,
            allowOrphans=True,
            keepWithNext=kind in ("title", "heading", "subheading"),
        )

    story = [
        Paragraph(escape(text).replace("\t", "    ").replace("\n", "<br/>"), styles[kind])
        for kind, text in paragraphs
    ]

    def page_furniture(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.setFillColor(colors.HexColor("#596773"))
        canvas.drawString(42, A4[1] - 27, "SAMRUK KAZYNA · Протокол совещания")
        canvas.drawRightString(A4[0] - 42, 25, f"Страница {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=42,
        rightMargin=42,
        topMargin=48,
        bottomMargin=45,
        title=_clean(title, "Протокол совещания"),
        author="SAMRUK KAZYNA",
        allowSplitting=True,
        pageCompression=1,
    )
    document.build(story, onFirstPage=page_furniture, onLaterPages=page_furniture)


def export_protocol(title: str, result: dict, directory: Path) -> None:
    """Write both protocol.docx and protocol.pdf under the caller's private job directory.

    The backend must advertise completed exports only after this function succeeds.
    Models, web services, a browser, and an installed PDF printer are not used.
    """
    directory = Path(directory)
    paragraphs = _paragraphs(title, result)
    fonts = _pdf_fonts()
    directory.mkdir(mode=0o700, parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".protocol-", dir=directory) as temporary:
        staging = Path(temporary)
        docx_path = staging / "protocol.docx"
        pdf_path = staging / "protocol.pdf"
        _write_docx(docx_path, title, paragraphs)
        _write_pdf(pdf_path, title, paragraphs, fonts)
        # Restrictive Unix permissions; on Windows the server directory's ACL applies.
        docx_path.chmod(0o600)
        pdf_path.chmod(0o600)
        # Both documents are fully written before either is made available to the API.
        os.replace(docx_path, directory / "protocol.docx")
        os.replace(pdf_path, directory / "protocol.pdf")
