"""Backend regressions without neural models, network calls, or user data.

Run from the repository root:
    python -m unittest discover -s tests -p "test_backend_api.py" -v

Each test redirects the storage directory into a temporary folder. Inference is
replaced with fixed fixtures; the export roundtrip uses the real local exporter.
"""

from __future__ import annotations

import asyncio
import io
import json
import queue
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi import HTTPException, UploadFile
from fastapi.testclient import TestClient

from backend.app import main
from backend.app.agent import orchestrator
from backend.app.agent.models import ProtocolResult, Segment, Task, Utterance


class BackendAPITests(unittest.TestCase):
    def setUp(self) -> None:
        scratch = Path(__file__).resolve().parents[1] / "work"
        scratch.mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(prefix="backend-api-test-", dir=scratch)
        self.addCleanup(self.temporary.cleanup)
        self.data = Path(self.temporary.name)
        storage_patch = patch.object(main, "DATA", self.data)
        storage_patch.start()
        self.addCleanup(storage_patch.stop)

    @staticmethod
    def fixture_result(meeting_id: str) -> ProtocolResult:
        return ProtocolResult(
            meeting_id=meeting_id,
            status="completed",
            summary=["Согласованы сроки подготовки отчёта."],
            tasks=[Task("Әлия", "2026-09-30", "Подготовить отчёт", "Отчёт к среде.")],
            transcript=[Utterance("SPEAKER_00", 0, 1, "Әлия, отчёт к среде.")],
            warnings=[],
        )

    def complete_fixture(self, audio_path, meeting_id, output_dir, on_progress=None):
        if on_progress is not None:
            on_progress("export", 90, "Создаём протоколы…")
        result = self.fixture_result(meeting_id)
        main.export_files(result, output_dir)
        return result

    def upload_completed_fixture(self) -> str:
        # Closing the app lifespan joins the worker deterministically. Tests do
        # not wait arbitrary durations and cannot touch a running user's server.
        with patch.object(main.agent, "process", side_effect=self.complete_fixture):
            with TestClient(main.app) as client:
                response = client.post(
                    "/api/meetings", files={"audio": ("Review.wav", b"fixture recording")}
                )
                self.assertEqual(response.status_code, 202, response.text)
                meeting_id = response.json()["meeting_id"]
        return meeting_id

    def test_upload_rejections_leave_no_partial_recordings(self) -> None:
        with patch.object(main.agent, "process") as inference:
            with TestClient(main.app) as client:
                for name, contents, expected in (
                    ("empty.wav", b"", 400),
                    ("program.exe", b"1234", 415),
                ):
                    with self.subTest(name=name):
                        response = client.post(
                            "/api/meetings", files={"audio": (name, contents)}
                        )
                        self.assertEqual(response.status_code, expected, response.text)
                        self.assertEqual(list(self.data.iterdir()), [])
                with patch.object(main, "MAX_UPLOAD_BYTES", 3):
                    response = client.post(
                        "/api/meetings", files={"audio": ("large.wav", b"1234")}
                    )
                    self.assertEqual(response.status_code, 413)
                self.assertEqual(list(self.data.iterdir()), [])
                inference.assert_not_called()

    def test_streaming_size_limit_applies_without_declared_size(self) -> None:
        # Multipart normally supplies a size. Exercise the byte-count guard too,
        # since the server must not trust only the declared upload metadata.
        recording = UploadFile(filename="recording.wav", file=io.BytesIO(b"1234"))
        with patch.object(main, "MAX_UPLOAD_BYTES", 3):
            with self.assertRaises(HTTPException) as caught:
                asyncio.run(main.create_meeting(recording))
        self.assertEqual(caught.exception.status_code, 413)
        self.assertTrue(recording.file.closed)
        self.assertEqual(list(self.data.iterdir()), [])

    def test_status_follows_real_orchestrator_stage_transitions(self) -> None:
        started: queue.Queue[str] = queue.Queue()
        release = threading.Semaphore(0)

        def wait_at(stage: str) -> None:
            started.put(stage)
            if not release.acquire(timeout=5):
                raise RuntimeError("Test did not release the controlled inference stage")

        def transcribe(_path):
            wait_at("transcribe")
            return [Segment(0, 1, "Әлия, отчёт к среде.")], []

        def diarize(_path):
            wait_at("diarize")
            return [{"start": 0, "end": 1, "speaker_id": "SPEAKER_01"}], []

        def analyze(_text):
            wait_at("analyze")
            fixture = self.fixture_result("unused")
            return fixture.tasks, fixture.summary, []

        real_export = main.export_files

        def export(result, folder):
            wait_at("export")
            return real_export(result, folder)

        with patch.object(orchestrator, "transcribe", side_effect=transcribe), \
             patch.object(orchestrator, "diarize", side_effect=diarize), \
             patch.object(orchestrator, "analyze", side_effect=analyze), \
             patch.object(orchestrator, "export_files", side_effect=export):
            with TestClient(main.app) as client:
                try:
                    response = client.post(
                        "/api/meetings", files={"audio": ("Review.MP4", b"fixture")}
                    )
                    self.assertEqual(response.status_code, 202, response.text)
                    accepted = response.json()
                    meeting_id = accepted["meeting_id"]
                    self.assertEqual(accepted["status"], "processing")
                    for stage, progress in (
                        ("transcribe", 10), ("diarize", 40), ("analyze", 65), ("export", 90)
                    ):
                        self.assertEqual(started.get(timeout=5), stage)
                        state = client.get(accepted["status_url"]).json()
                        self.assertEqual(state["status"], "processing")
                        self.assertEqual((state["stage"], state["progress"]), (stage, progress))
                        self.assertEqual(client.get(accepted["result_url"]).status_code, 409)
                        self.assertEqual(client.get("/health").json()["status"], "ok")
                        release.release()
                finally:
                    # Avoid hanging the lifespan join if an assertion fails.
                    for _ in range(4):
                        release.release()
        with TestClient(main.app) as client:
            state = client.get(f"/api/meetings/{meeting_id}/status").json()
            self.assertEqual((state["status"], state["progress"]), ("completed", 100))
            result = client.get(f"/api/meetings/{meeting_id}").json()
            self.assertEqual(result["transcript"][0]["speaker"], "SPEAKER_01")
            self.assertEqual(result["revision"], 1)

    def test_task_revisions_metadata_and_export_roundtrip(self) -> None:
        meeting_id = self.upload_completed_fixture()
        with TestClient(main.app) as client:
            original = client.get(f"/api/meetings/{meeting_id}").json()
            changes = {
                "revision": 1,
                "tasks": [{"description": "Подготовить финальный отчёт", "assignee": "Әлия",
                           "deadline": "2026-09-30", "status": "Выполнено"}],
            }
            response = client.put(f"/api/meetings/{meeting_id}/tasks", json=changes)
            self.assertEqual(response.status_code, 200, response.text)
            updated = response.json()
            self.assertEqual(updated["revision"], 2)
            for field in ("title", "created_at"):
                self.assertEqual(updated[field], original[field])
            self.assertEqual(
                client.put(f"/api/meetings/{meeting_id}/tasks", json=changes).status_code, 409
            )
            saved = client.get(f"/api/meetings/{meeting_id}").json()
            self.assertEqual(saved, updated)
            self.assertEqual(client.get("/api/meetings").json()["meetings"], [updated])

            docx = client.get(f"/api/meetings/{meeting_id}/download/docx")
            self.assertEqual(docx.status_code, 200)
            document = Document(io.BytesIO(docx.content))
            task_row = [cell.text for cell in document.tables[0].rows[1].cells]
            self.assertEqual(task_row[:4], ["Подготовить финальный отчёт", "Әлия", "2026-09-30", "Выполнено"])
            self.assertIn("Әлия, отчёт к среде.", " ".join(p.text for p in document.paragraphs))

            pdf = client.get(f"/api/meetings/{meeting_id}/download/pdf")
            self.assertEqual(pdf.status_code, 200)
            self.assertTrue(pdf.content.startswith(b"%PDF-"))
            self.assertTrue(pdf.content.rstrip().endswith(b"%%EOF"))
            self.assertIn(b"/ToUnicode", pdf.content)  # Embedded Cyrillic/Kazakh font mapping.

    def test_export_failure_does_not_replace_saved_revision(self) -> None:
        meeting_id = self.upload_completed_fixture()
        with TestClient(main.app) as client:
            original = client.get(f"/api/meetings/{meeting_id}").json()
            original_pdf = client.get(f"/api/meetings/{meeting_id}/download/pdf").content
            with patch.object(main, "export_files", side_effect=RuntimeError("fixture failure")):
                with self.assertLogs(main.logger.name, level="ERROR"):
                    response = client.put(
                        f"/api/meetings/{meeting_id}/tasks", json={"revision": 1, "tasks": []}
                    )
            self.assertEqual(response.status_code, 500)
            self.assertEqual(client.get(f"/api/meetings/{meeting_id}").json(), original)
            self.assertEqual(
                client.get(f"/api/meetings/{meeting_id}/download/pdf").content, original_pdf
            )

    def test_failed_job_remains_available_after_server_restart(self) -> None:
        def fail(_path, _meeting_id, _folder, on_progress=None):
            on_progress("analyze", 65, "Извлекаем поручения…")
            raise RuntimeError("controlled inference failure")

        with patch.object(main.agent, "process", side_effect=fail):
            with self.assertLogs(main.logger.name, level="ERROR"):
                with TestClient(main.app) as client:
                    response = client.post(
                        "/api/meetings", files={"audio": ("failure.wav", b"fixture")}
                    )
                    self.assertEqual(response.status_code, 202)
                    meeting_id = response.json()["meeting_id"]
        with TestClient(main.app) as client:
            response = client.get(f"/api/meetings/{meeting_id}/status")
            self.assertEqual(response.status_code, 200)
            state = response.json()
            self.assertEqual((state["status"], state["failed_stage"]), ("failed", "analyze"))
            self.assertEqual(state["progress"], 65)
            self.assertEqual(client.get(f"/api/meetings/{meeting_id}").status_code, 422)
            self.assertEqual(client.get("/api/meetings").json()["meetings"], [])

    def test_restart_recovers_interrupted_jobs_and_rejects_unknown_ids(self) -> None:
        meeting_id = "a" * 32
        folder = self.data / meeting_id
        folder.mkdir()
        (folder / "status.json").write_text(json.dumps({
            "meeting_id": meeting_id, "status": "processing", "stage": "diarize",
            "progress": 40, "message": "pending",
        }), encoding="utf-8")
        with TestClient(main.app) as client:
            state = client.get(f"/api/meetings/{meeting_id}/status").json()
            self.assertEqual(state["status"], "failed")
            for invalid_id in ("not-an-id", "b" * 32):
                with self.subTest(meeting_id=invalid_id):
                    self.assertEqual(client.get(f"/api/meetings/{invalid_id}/status").status_code, 404)
            self.assertEqual(client.get(f"/api/meetings/{meeting_id}/download/exe").status_code, 400)


if __name__ == "__main__":
    unittest.main()
